import sys
import glob
import random
import argparse
import joblib
import itertools
import json
import numpy as np
import pandas as pd
import seaborn as sns
from pathlib import Path
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from tslearn.clustering import TimeSeriesKMeans

from sklearn.model_selection import ParameterGrid, StratifiedKFold, cross_validate, train_test_split
from sklearn.ensemble import RandomForestClassifier, BaggingClassifier, VotingClassifier
from sklearn.svm import SVC
from sklearn.naive_bayes import GaussianNB
from sklearn.tree import DecisionTreeClassifier
from sklearn.neural_network import MLPClassifier
from sklearn.linear_model import LogisticRegression
from tabpfn import TabPFNClassifier
from sklearn.metrics import (
    confusion_matrix, accuracy_score, classification_report,
    precision_score, recall_score, f1_score, fbeta_score, make_scorer, roc_auc_score,
)

sys.path.insert(0, str(Path(__file__).parent.parent.parent.parent))

from config.config import SEED, PARAM_GRID, PATH_MODELS, PATH_DATASET

# ──────────────────────────────────────────────────────────────────────────────
# Constants
# ──────────────────────────────────────────────────────────────────────────────

classmap = {"no_FHB": 0, "FHB": 1}

PATH_DATA   = Path(PATH_DATASET + "Oslo_viability/markers/")
PATH_MODELS = Path(PATH_MODELS  + "assessment/vIAbility/markers/ML/")
PATH_OUTPUT = Path("output/assessment/vIAbility/markers/ML/")

N_CLUSTERS = 5
CV_FOLDS   = 5
VAL_SIZE   = 0.2

random.seed(SEED)
np.random.seed(SEED)

CLASSIFIERS = {
    "RandomForest":      RandomForestClassifier,
    "SVC":               SVC,
    "GaussianNB":        GaussianNB,
    "DecisionTree":      DecisionTreeClassifier,
    "MLPClassifier":     MLPClassifier,
    "LogisticRegression": LogisticRegression,
    "TabPFN":            TabPFNClassifier,
}

CV_SCORING = {
    "accuracy":  "accuracy",
    "precision": make_scorer(precision_score, pos_label=1, zero_division=0),
    "recall":    make_scorer(recall_score,    pos_label=1, zero_division=0),
    "fbeta":     make_scorer(fbeta_score, beta=0.5, pos_label=1, zero_division=0),
    "auroc":     make_scorer(roc_auc_score, needs_proba=True),
}

# ──────────────────────────────────────────────────────────────────────────────
# Data loading
# ──────────────────────────────────────────────────────────────────────────────

def load_data(path: Path, split: str = "train"):
    labels = []
    X = None
    files = sorted(glob.glob(str(path / split / "*/*.csv")))

    if not files:
        raise ValueError(f"No CSV files found in {path / split}")

    for i, file in enumerate(files):
        df = pd.read_csv(file)
        df.replace([np.inf, -np.inf], np.nan, inplace=True)
        df.fillna(0, inplace=True)
        for col in df.columns:
            if df[col].dtype == object:
                df[col] = pd.Categorical(df[col]).codes
        columns = df.columns[2:-1]
        if i == 0:
            X = {c: [] for c in columns}
        for c in columns:
            X[c].append(df[c].tolist())
        labels.append(classmap[Path(file).parent.stem])

    return {c: np.array(X[c]) for c in columns}, np.array(labels, dtype=int)

# ──────────────────────────────────────────────────────────────────────────────
# Time-series clustering
# ──────────────────────────────────────────────────────────────────────────────

def fit_cluster_models(X_train: dict, columns, n_clusters: int, seed: int):
    models = {}
    for v in columns:
        ts = X_train[v].reshape(X_train[v].shape[0], X_train[v].shape[1], 1)
        models[v] = TimeSeriesKMeans(
            n_clusters=n_clusters, metric="euclidean", max_iter=100, random_state=seed,
        ).fit(ts)
        print(f"  Clustered feature '{v}'")
    return models


def apply_clusters(cluster_models: dict, X: dict, n_samples: int) -> np.ndarray:
    parts = []
    for v, km in cluster_models.items():
        ts  = X[v].reshape(X[v].shape[0], X[v].shape[1], 1)
        ids = km.predict(ts).reshape(n_samples, 1)
        parts.append(ids)
    return np.concatenate(parts, axis=1)

# ──────────────────────────────────────────────────────────────────────────────
# Hyperparameter search
# ──────────────────────────────────────────────────────────────────────────────

def study_hyperparameters(model_cls, param_grid, X_train, y_train, X_val, y_val):
    best_fbeta  = -1
    best_params = None

    for params in ParameterGrid(param_grid):
        try:
            clf   = model_cls(**params)
            clf.fit(X_train, y_train)
            preds = clf.predict(X_val)
            acc   = accuracy_score(y_val, preds)
            fbeta = fbeta_score(y_val, preds, beta=0.5, pos_label=1, zero_division=0)
            print(f"  {params} → val_acc={acc:.4f} val_fbeta={fbeta:.4f}")
            if fbeta > best_fbeta:
                best_fbeta  = fbeta
                best_params = params
        except Exception as e:
            print(f"  FAILED {params} → {e}")

    print(f"\n  Best params: {best_params}")
    print(f"  Best validation Fbeta: {best_fbeta:.4f}\n")
    return best_params

# ──────────────────────────────────────────────────────────────────────────────
# Ensemble utilities
# ──────────────────────────────────────────────────────────────────────────────

def build_bagging_classifiers(best_params: dict, seed: int):
    bagging_models = {}
    for name, clf_cls in CLASSIFIERS.items():
        base = clf_cls(**best_params[name])
        bagging_models[f"Bagging_{name}"] = BaggingClassifier(
            estimator=base, n_estimators=10, random_state=seed, n_jobs=-1,
        )
    return bagging_models


def _make_estimator(name: str, best_params: dict):
    params = dict(best_params[name])
    if name == "SVC":
        params["probability"] = True
    return CLASSIFIERS[name](**params)


def search_voting_combinations(
    best_params: dict,
    X_train: np.ndarray,
    y_train: np.ndarray,
    X_val: np.ndarray,
    y_val: np.ndarray,
    min_size: int = 2,
    max_size: int = None,
    voting: str = "both",
    top_k: int = 3,
):
    names      = list(CLASSIFIERS.keys())
    max_size   = max_size or len(names)
    strategies = ["hard", "soft"] if voting == "both" else [voting]
    results    = {s: [] for s in strategies}

    total_combos = sum(
        len(list(itertools.combinations(names, r)))
        for r in range(min_size, max_size + 1)
    )
    print(f"  Evaluating {total_combos} combinations (sizes {min_size}–{max_size})")

    for size in range(min_size, max_size + 1):
        for combo in itertools.combinations(names, size):
            estimators = [(n, _make_estimator(n, best_params)) for n in combo]
            combo_tag  = "+".join(combo)
            for strategy in strategies:
                try:
                    vc = VotingClassifier(estimators=estimators, voting=strategy, n_jobs=-1)
                    vc.fit(X_train, y_train)
                    preds = vc.predict(X_val)
                    acc   = accuracy_score(y_val, preds)
                    fbeta = fbeta_score(y_val, preds, beta=0.5, pos_label=1, zero_division=0)
                    results[strategy].append((fbeta, combo_tag, estimators, strategy))
                    print(f"    [{strategy}] {combo_tag:<60} val_acc={acc:.4f} val_fbeta={fbeta:.4f}")
                except Exception as e:
                    print(f"    [{strategy}] {combo_tag} skipped ({e})")

    best_voters = {}
    for strategy, entries in results.items():
        entries.sort(key=lambda x: x[0], reverse=True)
        print(f"\nTop-{top_k} [{strategy}]:")
        for rank, (fbeta, tag, estimators, strategy_name) in enumerate(entries[:top_k], 1):
            key = f"Voting_{strategy}_{tag}"
            best_voters[key] = {"estimators": estimators, "strategy": strategy_name}
            print(f"    #{rank} {tag} val_fbeta={fbeta:.4f}")

    return best_voters

# ──────────────────────────────────────────────────────────────────────────────
# Cross-validation
# ──────────────────────────────────────────────────────────────────────────────

def cross_validate_models(
    models: dict,
    X_train: np.ndarray,
    y_train: np.ndarray,
    n_folds: int = CV_FOLDS,
    seed: int = SEED,
):
    cv = StratifiedKFold(n_splits=n_folds, shuffle=True, random_state=seed)
    cv_results = {}

    print("\n" + "=" * 70)
    print(f"Cross-validation ({n_folds}-fold)")
    print("=" * 70)

    for name, clf in models.items():
        print(f"\n[{name}]")
        try:
            scores = cross_validate(
                clf, X_train, y_train,
                cv=cv, scoring=CV_SCORING, n_jobs=-1, return_train_score=False,
            )
            summary = {}
            for metric in CV_SCORING:
                vals = scores[f"test_{metric}"]
                summary[metric] = {
                    "mean": float(vals.mean()),
                    "std":  float(vals.std()),
                    "per_fold": vals.tolist(),
                }
                print(f"  {metric:<10} mean={vals.mean()*100:.2f}% std={vals.std()*100:.2f}%")
            cv_results[name] = summary
        except Exception as e:
            print(f"  CV failed → {e}")
            cv_results[name] = None

    return cv_results

# ──────────────────────────────────────────────────────────────────────────────
# Consistency report
# ──────────────────────────────────────────────────────────────────────────────

def _consistency_flag(cv_mean: float, test_score: float, cv_std: float):
    diff      = test_score - cv_mean
    threshold = max(1.5 * cv_std, 0.02)
    if diff > threshold:
        return "⚠ optimistic"
    elif diff < -threshold:
        return "⚠ pessimistic"
    return "✓ consistent"


def print_consistency_report(cv_results: dict, test_metrics: dict, out_dir: Path):
    lines  = []
    header = "\n" + "=" * 90 + "\nConsistency report (CV vs Test)\n" + "=" * 90
    print(header)
    lines.append(header)

    for name in cv_results:
        if cv_results[name] is None or name not in test_metrics:
            continue
        block = [f"\n{name}"]
        for metric in CV_SCORING:
            cv_mean    = cv_results[name][metric]["mean"]
            cv_std     = cv_results[name][metric]["std"]
            test_score = test_metrics[name][metric]
            flag       = _consistency_flag(cv_mean, test_score, cv_std)
            row = (
                f"  {metric:<10} "
                f"CV={cv_mean*100:.2f}% ± {cv_std*100:.2f}% | "
                f"TEST={test_score*100:.2f}% | {flag}"
            )
            block.append(row)
        block_str = "\n".join(block)
        print(block_str)
        lines.append(block_str)

    report_path = out_dir / "cv_consistency_report.txt"
    report_path.write_text("\n".join(lines))
    print(f"\nSaved → {report_path}")

# ──────────────────────────────────────────────────────────────────────────────
# Plot CV vs Test
# ──────────────────────────────────────────────────────────────────────────────

def plot_cv_results(cv_results: dict, test_metrics: dict, out_dir: Path, metric: str = "fbeta"):
    names       = []
    cv_means    = []
    cv_stds     = []
    test_scores = []

    for name, res in cv_results.items():
        if res is None or name not in test_metrics:
            continue
        names.append(name)
        cv_means.append(res[metric]["mean"] * 100)
        cv_stds.append(res[metric]["std"]   * 100)
        test_scores.append(test_metrics[name][metric] * 100)

    if not names:
        return

    x     = np.arange(len(names))
    width = 0.35
    fig, ax = plt.subplots(figsize=(max(12, len(names) * 1.3), 6))

    ax.bar(x - width / 2, cv_means,    width, yerr=cv_stds, capsize=4,
           label="CV mean ± std", color="steelblue", alpha=0.85)
    ax.bar(x + width / 2, test_scores, width,
           label="Test", color="coral", alpha=0.85)

    ax.set_ylabel(f"{metric.upper()} (%)")
    ax.set_title(f"CV vs Test ({metric})")
    ax.set_xticks(x)
    ax.set_xticklabels(names, rotation=40, ha="right", fontsize=8)
    ax.legend()
    ax.set_ylim(0, 110)
    fig.tight_layout()

    out_path = out_dir / f"cv_vs_test_{metric}.png"
    fig.savefig(out_path, dpi=150)
    plt.close(fig)
    print(f"Saved → {out_path}")

# ──────────────────────────────────────────────────────────────────────────────
# Save CV results
# ──────────────────────────────────────────────────────────────────────────────

def save_cv_results(cv_results: dict, out_dir: Path, test_metrics: dict = None, best_params: dict = None):
    out_dir.mkdir(parents=True, exist_ok=True)

    # ── JSON (unchanged) ──────────────────────────────────────────────────────
    json_path = out_dir / "cv_results.json"
    with open(json_path, "w") as f:
        json.dump(cv_results, f, indent=4)
    print(f"Saved CV results → {json_path}")

    # ── Excel ─────────────────────────────────────────────────────────────────
    wb = Workbook()

    # ── Sheet 1: CV summary ───────────────────────────────────────────────────
    ws_cv = wb.active
    ws_cv.title = "CV Results"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    center      = Alignment(horizontal="center")

    cv_metrics  = list(CV_SCORING.keys())
    cv_headers  = ["Model"] + [f"{m}_mean" for m in cv_metrics] + [f"{m}_std" for m in cv_metrics]
    ws_cv.append(cv_headers)
    for col, _ in enumerate(cv_headers, 1):
        cell = ws_cv.cell(1, col)
        cell.fill, cell.font, cell.alignment = header_fill, header_font, center

    for model_name, metrics in cv_results.items():
        if metrics is None:
            continue
        row = [model_name]
        for m in cv_metrics:
            row.append(round(metrics[m]["mean"], 4))
        for m in cv_metrics:
            row.append(round(metrics[m]["std"], 4))
        ws_cv.append(row)

    for col in range(1, len(cv_headers) + 1):
        ws_cv.column_dimensions[get_column_letter(col)].width = 18

    # ── Sheet 2: Test metrics ─────────────────────────────────────────────────
    if test_metrics:
        ws_test = wb.create_sheet("Test Metrics")
        test_metric_keys = list(next(iter(test_metrics.values())).keys())
        t_headers = ["Model"] + test_metric_keys
        ws_test.append(t_headers)
        for col, _ in enumerate(t_headers, 1):
            cell = ws_test.cell(1, col)
            cell.fill, cell.font, cell.alignment = header_fill, header_font, center
        for model_name, metrics in test_metrics.items():
            row = [model_name] + [round(metrics[k], 4) for k in test_metric_keys]
            ws_test.append(row)
        for col in range(1, len(t_headers) + 1):
            ws_test.column_dimensions[get_column_letter(col)].width = 18

    # ── Sheet 3: Best parameters ──────────────────────────────────────────────
    if best_params:
        ws_params = wb.create_sheet("Best Parameters")
        # Collect all param keys across base classifiers
        all_param_keys = sorted({k for p in best_params.values() if p for k in p.keys()})
        p_headers = ["Model"] + all_param_keys
        ws_params.append(p_headers)
        for col, _ in enumerate(p_headers, 1):
            cell = ws_params.cell(1, col)
            cell.fill, cell.font, cell.alignment = header_fill, header_font, center
        for model_name, params in best_params.items():
            if params is None:
                row = [model_name] + ["N/A"] * len(all_param_keys)
            else:
                row = [model_name] + [params.get(k, "") for k in all_param_keys]
            row = [str(item) if isinstance(item, tuple) else item for item in row]
            ws_params.append(row)
        for col in range(1, len(p_headers) + 1):
            ws_params.column_dimensions[get_column_letter(col)].width = 22

    xlsx_path = out_dir / "cv_results.xlsx"
    wb.save(xlsx_path)
    print(f"Saved CV results → {xlsx_path}")

# ──────────────────────────────────────────────────────────────────────────────
# Evaluation
# ──────────────────────────────────────────────────────────────────────────────

def _get_proba(clf, X):
    """Return positive-class probability scores, falling back to decision_function."""
    if hasattr(clf, "predict_proba"):
        return clf.predict_proba(X)[:, 1]
    if hasattr(clf, "decision_function"):
        return clf.decision_function(X)
    return None

def evaluate_and_save(models: dict, X_test: np.ndarray, y_test: np.ndarray, path_model: str, out_dir: str):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    class_names   = ["no_FHB", "FHB"]
    summary_lines = []
    test_metrics  = {}

    for name, clf in models.items():
        print("\n" + "=" * 80)
        print(name)
        print("=" * 80)

        y_pred = clf.predict(X_test)
        acc    = accuracy_score(y_test, y_pred)
        prec   = precision_score(y_test, y_pred, pos_label=1, zero_division=0)
        rec    = recall_score(y_test, y_pred, pos_label=1, zero_division=0)
        fbeta  = fbeta_score(y_test, y_pred, beta=0.5, pos_label=1, zero_division=0)
        proba  = _get_proba(clf, X_test)
        auroc  = roc_auc_score(y_test, proba) if proba is not None else float("nan")

        test_metrics[name] = {"accuracy": acc, "precision": prec, "recall": rec, "fbeta": fbeta, "auroc": auroc}

        report = classification_report(y_test, y_pred, target_names=class_names, zero_division=0)

        result_str = (
            f"{name}\n\n"
            f"Accuracy : {acc:.4f}\n"
            f"Precision: {prec:.4f}\n"
            f"Recall   : {rec:.4f}\n"
            f"Fbeta(0.5): {fbeta:.4f}\n"
            f"AUROC    : {auroc:.4f}\n\n"
            f"{report}"
        )
        print(result_str)
        (out_dir / f"test_results_{name}.txt").write_text(result_str)

        summary_lines.append(
            f"{name:<50} acc={acc:.4f} prec={prec:.4f} rec={rec:.4f} fbeta={fbeta:.4f} auroc={auroc:.4f}"
        )

        # Confusion matrix
        cm      = confusion_matrix(y_test, y_pred)
        fig, ax = plt.subplots(figsize=(5, 4))
        sns.heatmap(cm, annot=True, fmt="d", cmap="Blues",
                    xticklabels=class_names, yticklabels=class_names, ax=ax)
        ax.set_xlabel("Predicted")
        ax.set_ylabel("True")
        ax.set_title(name)
        fig.tight_layout()
        fig.savefig(out_dir / f"confusion_matrix_{name}.png", dpi=150)
        plt.close(fig)

        # Save model
        model_path = Path(path_model) / f"{name}.joblib"
        model_path.parent.mkdir(parents=True, exist_ok=True)
        joblib.dump(clf, model_path)
        print(f"Saved model → {model_path}")

    summary = "\n".join(summary_lines)
    (out_dir / "all_results.txt").write_text(summary)
    print("\nFinal summary")
    print(summary)
    return test_metrics

# ──────────────────────────────────────────────────────────────────────────────
# Main
# ──────────────────────────────────────────────────────────────────────────────

def main(args):
    print("Loading data...")
    X_full_train, y_full_train = load_data(PATH_DATA, split="train")
    X_test, y_test             = load_data(PATH_DATA, split="test")
    columns                    = list(X_full_train.keys())
    print(f"Features: {columns}")

    # Clustering
    print("\nFitting clustering models...")
    cluster_models  = fit_cluster_models(X_full_train, columns, args.n_clusters, SEED)
    X_full_train_cl = apply_clusters(cluster_models, X_full_train, len(y_full_train))
    X_test_cl       = apply_clusters(cluster_models, X_test, len(y_test))
    print(f"\nClustered shape: {X_full_train_cl.shape}")

    # Validation split
    X_train_cl, X_val_cl, y_train, y_val = train_test_split(
        X_full_train_cl, y_full_train,
        test_size=args.val_size, stratify=y_full_train, random_state=SEED,
    )
    print("\nDataset split:")
    print(f"Train: {len(y_train)}")
    print(f"Val  : {len(y_val)}")
    print(f"Test : {len(y_test)}")

    # Hyperparameter search
    print("\nSearching hyperparameters...")
    best_params = {}
    for name, clf_cls in CLASSIFIERS.items():
        print(f"\n--- {name} ---")
        best_params[name] = study_hyperparameters(
            clf_cls, PARAM_GRID[name], X_train_cl, y_train, X_val_cl, y_val,
        )

    # Refit on train + validation
    X_final_train = np.concatenate([X_train_cl, X_val_cl], axis=0)
    y_final_train = np.concatenate([y_train, y_val], axis=0)

    # Base models
    print("\nTraining final base models...")
    best_models = {}
    for name in CLASSIFIERS:
        clf = CLASSIFIERS[name](**best_params[name])
        clf.fit(X_final_train, y_final_train)
        best_models[name] = clf

    # Bagging
    print("\nTraining bagging models...")
    bagging_models = build_bagging_classifiers(best_params, SEED)
    for name, clf in bagging_models.items():
        print(f"  Fitting {name}")
        clf.fit(X_final_train, y_final_train)

    # Voting search
    print("\nSearching voting combinations...")
    voting_configs = search_voting_combinations(
        best_params, X_train_cl, y_train, X_val_cl, y_val,
        min_size=args.vote_min, max_size=args.vote_max,
        voting=args.voting, top_k=args.vote_top_k,
    )

    # Final voting models
    final_voting_models = {}
    print("\nRetraining selected voting models...")
    for name, cfg in voting_configs.items():
        clf = VotingClassifier(estimators=cfg["estimators"], voting=cfg["strategy"], n_jobs=-1)
        clf.fit(X_final_train, y_final_train)
        final_voting_models[name] = clf

    all_models = {**best_models, **bagging_models, **final_voting_models}

    # Cross-validation
    cv_results = cross_validate_models(all_models, X_train_cl, y_train, n_folds=args.cv_folds, seed=SEED)

    # Final test evaluation
    test_metrics = evaluate_and_save(
        all_models, X_test_cl, y_test,
        path_model=args.path_models, out_dir=args.out_dir,
    )

    save_cv_results(cv_results, Path(args.out_dir), test_metrics=test_metrics, best_params=best_params)

    # CV vs Test comparison
    print_consistency_report(cv_results, test_metrics, Path(args.out_dir))
    for metric in CV_SCORING:
        plot_cv_results(cv_results, test_metrics, Path(args.out_dir), metric=metric)

# ──────────────────────────────────────────────────────────────────────────────
# Entry point
# ──────────────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Viability classification via time-series clustering + ML"
    )
    parser.add_argument("--n_clusters",  type=int,   default=N_CLUSTERS)
    parser.add_argument("--path_models",             default=PATH_MODELS)
    parser.add_argument("--out_dir",                 default=PATH_OUTPUT)
    parser.add_argument("--vote_min",    type=int,   default=2)
    parser.add_argument("--vote_max",    type=int,   default=None)
    parser.add_argument("--voting",                  default="both", choices=["hard", "soft", "both"])
    parser.add_argument("--vote_top_k",  type=int,   default=3)
    parser.add_argument("--cv_folds",    type=int,   default=CV_FOLDS)
    parser.add_argument("--val_size",    type=float, default=VAL_SIZE)
    main(parser.parse_args())