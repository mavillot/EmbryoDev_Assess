import sys, glob, random, argparse, joblib, itertools, json, csv
from datetime import datetime
from pathlib import Path
import numpy as np
import pandas as pd
import seaborn as sns
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from tslearn.clustering import TimeSeriesKMeans
from sklearn.model_selection import ParameterGrid, StratifiedKFold, cross_validate, train_test_split
from sklearn.ensemble import RandomForestClassifier, BaggingClassifier, VotingClassifier
from sklearn.svm import SVC
from sklearn.naive_bayes import GaussianNB
from sklearn.tree import DecisionTreeClassifier
from sklearn.neural_network import MLPClassifier
from sklearn.linear_model import LogisticRegression
from tabpfn import TabPFNClassifier
from sklearn.metrics import confusion_matrix, accuracy_score, classification_report, precision_score, recall_score, f1_score, fbeta_score, make_scorer, roc_auc_score

sys.path.insert(0, str(Path(__file__).parent.parent.parent.parent))
from config.config import SEED, PARAM_GRID, PATH_MODELS, PATH_DATASET

classmap = {"no_FHB": 0, "FHB": 1}
PATH_DATA = Path(PATH_DATASET + "Oslo_viability/markers/")
PATH_MODELS = Path(PATH_MODELS + "assessment/vIAbility/markers/ML/")
PATH_OUTPUT = Path("results/assessment/vIAbility/features/markers/ML/")
N_CLUSTERS, CV_FOLDS, VAL_SIZE = 5, 5, 0.2
METRICS_CSV = str(PATH_OUTPUT / "experiment_results.csv")

def append_metrics_row(csv_path: str, row: dict):
    p = Path(csv_path)
    p.parent.mkdir(parents=True, exist_ok=True)
    ex = p.exists()
    with open(p, "a", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(row.keys()))
        if not ex: writer.writeheader()
        writer.writerow(row)

CLASSIFIERS = {
    "RandomForest": RandomForestClassifier, "SVC": SVC, "GaussianNB": GaussianNB,
    "DecisionTree": DecisionTreeClassifier, "MLPClassifier": MLPClassifier,
    "LogisticRegression": LogisticRegression, "TabPFN": TabPFNClassifier
}
CV_SCORING = {
    "accuracy": "accuracy",
    "precision": make_scorer(precision_score, pos_label=1, zero_division=0),
    "recall": make_scorer(recall_score, pos_label=1, zero_division=0),
    "fbeta": make_scorer(fbeta_score, beta=0.5, pos_label=1, zero_division=0),
    "f1": make_scorer(f1_score, pos_label=1, zero_division=0),
    "auroc": None,
}


def _auroc_scorer(estimator, X, y):
    if hasattr(estimator, "predict_proba"):
        scores = estimator.predict_proba(X)[:, 1]
    elif hasattr(estimator, "decision_function"):
        scores = estimator.decision_function(X)
    else:
        raise AttributeError(f"{estimator.__class__.__name__} has neither predict_proba nor decision_function")
    return roc_auc_score(y, scores)


CV_SCORING["auroc"] = _auroc_scorer

def load_data(path: Path, split: str = "train"):
    files = sorted(glob.glob(str(path / split / "*/*.csv")))
    if not files: raise ValueError(f"No CSV files found in {path / split}")
    labels, X = [], None
    for i, file in enumerate(files):
        df = pd.read_csv(file).replace([np.inf, -np.inf], np.nan).fillna(0)
        for col in df.columns:
            if df[col].dtype == object: df[col] = pd.Categorical(df[col]).codes
        cols = df.columns[2:-1]
        if i == 0: X = {c: [] for c in cols}
        for c in cols: X[c].append(df[c].tolist())
        labels.append(classmap[Path(file).parent.stem])
    return {c: np.array(X[c]) for c in cols}, np.array(labels, dtype=int)

def fit_cluster_models(X_train: dict, columns, n_clusters: int, seed: int):
    models = {}
    for v in columns:
        ts = X_train[v].reshape(X_train[v].shape[0], X_train[v].shape[1], 1)
        models[v] = TimeSeriesKMeans(n_clusters=n_clusters, metric="euclidean", max_iter=100, random_state=seed).fit(ts)
    return models

def apply_clusters(cluster_models: dict, X: dict, n_samples: int) -> np.ndarray:
    return np.concatenate([km.predict(X[v].reshape(X[v].shape[0], X[v].shape[1], 1)).reshape(n_samples, 1) for v, km in cluster_models.items()], axis=1)

def study_hyperparameters(model_cls, param_grid, X_train, y_train, X_val, y_val):
    best_fbeta, best_params = -1, None
    for params in ParameterGrid(param_grid):
        try:
            clf = model_cls(**params).fit(X_train, y_train)
            fbeta = fbeta_score(y_val, clf.predict(X_val), beta=0.5, pos_label=1, zero_division=0)
            if fbeta > best_fbeta: best_fbeta, best_params = fbeta, params
        except Exception as e: print(f"  FAILED {params} → {e}")
    return best_params

def build_bagging_classifiers(best_params: dict, seed: int):
    return {f"Bagging_{n}": BaggingClassifier(estimator=cls(**best_params[n]), n_estimators=10, random_state=seed, n_jobs=-1) for n, cls in CLASSIFIERS.items()}

def _make_estimator(name: str, best_params: dict):
    p = dict(best_params[name])
    if name == "SVC": p["probability"] = True
    return CLASSIFIERS[name](**p)

def search_voting_combinations(best_params: dict, X_train, y_train, X_val, y_val, min_size=2, max_size=None, voting="both", top_k=3):
    names = list(CLASSIFIERS.keys())
    strats = ["hard", "soft"] if voting == "both" else [voting]
    results = {s: [] for s in strats}
    for size in range(min_size, (max_size or len(names)) + 1):
        for combo in itertools.combinations(names, size):
            estimators = [(n, _make_estimator(n, best_params)) for n in combo]
            for s in strats:
                try:
                    vc = VotingClassifier(estimators=estimators, voting=s, n_jobs=-1).fit(X_train, y_train)
                    fbeta = fbeta_score(y_val, vc.predict(X_val), beta=0.5, pos_label=1, zero_division=0)
                    results[s].append((fbeta, "+".join(combo), estimators, s))
                except: pass
    best_voters = {}
    for s, entries in results.items():
        entries.sort(key=lambda x: x[0], reverse=True)
        for fbeta, tag, ests, s_name in entries[:top_k]:
            best_voters[f"Voting_{s}_{tag}"] = {"estimators": ests, "strategy": s_name}
    return best_voters

def cross_validate_models(models: dict, X_train, y_train, n_folds=CV_FOLDS, seed=SEED):
    cv = StratifiedKFold(n_splits=n_folds, shuffle=True, random_state=seed)
    cv_results = {}
    for name, clf in models.items():
        try:
            scores = cross_validate(clf, X_train, y_train, cv=cv, scoring=CV_SCORING, n_jobs=-1)
            cv_results[name] = {m: {"mean": float(scores[f"test_{m}"].mean()), "std": float(scores[f"test_{m}"].std()), "per_fold": scores[f"test_{m}"].tolist()} for m in CV_SCORING}
        except: cv_results[name] = None
    return cv_results

def print_consistency_report(cv_results: dict, test_metrics: dict, out_dir: Path):
    lines = ["\n" + "="*90 + "\nConsistency report (CV vs Test)\n" + "="*90]
    for name, res in cv_results.items():
        if not res or name not in test_metrics: continue
        lines.append(f"\n{name}")
        for m in CV_SCORING:
            mean, std, ts = res[m]["mean"], res[m]["std"], test_metrics[name][m]
            diff, th = ts - mean, max(1.5 * std, 0.02)
            flag = "⚠ optimistic" if diff > th else ("⚠ pessimistic" if diff < -th else "✓ consistent")
            lines.append(f"  {m:<10} CV={mean*100:.2f}% ± {std*100:.2f}% | TEST={ts*100:.2f}% | {flag}")
    report_str = "\n".join(lines)
    print(report_str)
    (out_dir / "cv_consistency_report.txt").write_text(report_str)

def plot_cv_results(cv_results: dict, test_metrics: dict, out_dir: Path, metric="fbeta"):
    valid = [(n, res[metric]["mean"]*100, res[metric]["std"]*100, test_metrics[n][metric]*100) for n, res in cv_results.items() if res and n in test_metrics]
    if not valid: return
    names, cv_m, cv_s, t_s = zip(*valid)
    x, width = np.arange(len(names)), 0.35
    fig, ax = plt.subplots(figsize=(max(12, len(names) * 1.3), 6))
    ax.bar(x - width/2, cv_m, width, yerr=cv_s, capsize=4, label="CV mean ± std", color="steelblue", alpha=0.85)
    ax.bar(x + width/2, t_s, width, label="Test", color="coral", alpha=0.85)
    ax.set(ylabel=f"{metric.upper()} (%)", title=f"CV vs Test ({metric})", xticks=x, ylim=(0, 110))
    ax.set_xticklabels(names, rotation=40, ha="right", fontsize=8)
    ax.legend()
    fig.tight_layout()
    fig.savefig(out_dir / f"cv_vs_test_{metric}.png", dpi=150)
    plt.close(fig)

def save_cv_results(cv_results: dict, out_dir: Path, test_metrics=None, best_params=None):
    out_dir.mkdir(parents=True, exist_ok=True)
    with open(out_dir / "cv_results.json", "w") as f: json.dump(cv_results, f, indent=4)
    wb = Workbook()
    fill, font, center = PatternFill("solid", fgColor="1F4E79"), Font(bold=True, color="FFFFFF"), Alignment(horizontal="center")
    
    def format_sheet(ws, headers):
        ws.append(headers)
        for col in range(1, len(headers) + 1):
            cell = ws.cell(1, col)
            cell.fill, cell.font, cell.alignment = fill, font, center
            ws.column_dimensions[get_column_letter(col)].width = 20

    ws_cv = wb.active
    ws_cv.title = "CV Results"
    cv_metrics = list(CV_SCORING.keys())
    format_sheet(ws_cv, ["Model"] + [f"{m}_mean" for m in cv_metrics] + [f"{m}_std" for m in cv_metrics])
    for n, metrics in cv_results.items():
        if metrics: ws_cv.append([n] + [round(metrics[m]["mean"], 4) for m in cv_metrics] + [round(metrics[m]["std"], 4) for m in cv_metrics])

    if test_metrics:
        ws_t = wb.create_sheet("Test Metrics")
        t_keys = list(next(iter(test_metrics.values())).keys())
        format_sheet(ws_t, ["Model"] + t_keys)
        for n, metrics in test_metrics.items(): ws_t.append([n] + [round(metrics[k], 4) for k in t_keys])

    if best_params:
        ws_p = wb.create_sheet("Best Parameters")
        p_keys = sorted({k for p in best_params.values() if p for k in p.keys()})
        format_sheet(ws_p, ["Model"] + p_keys)
        for n, p in best_params.items():
            row = [n] + (["N/A"] * len(p_keys) if p is None else [p.get(k, "") for k in p_keys])
            ws_p.append([str(i) if isinstance(i, tuple) else i for i in row])

    wb.save(out_dir / "cv_results.xlsx")

def evaluate_and_save(models: dict, X_test, y_test, path_model: str, out_dir: str):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    test_metrics, summary_lines = {}, []
    for name, clf in models.items():
        y_pred = clf.predict(X_test)
        proba = clf.predict_proba(X_test)[:, 1] if hasattr(clf, "predict_proba") else (clf.decision_function(X_test) if hasattr(clf, "decision_function") else None)
        metrics = {
            "accuracy": accuracy_score(y_test, y_pred),
            "precision": precision_score(y_test, y_pred, pos_label=1, zero_division=0),
            "recall": recall_score(y_test, y_pred, pos_label=1, zero_division=0),
            "fbeta": fbeta_score(y_test, y_pred, beta=0.5, pos_label=1, zero_division=0),
            "f1": f1_score(y_test, y_pred, pos_label=1, zero_division=0),
            "auroc": roc_auc_score(y_test, proba) if proba is not None else float("nan")
        }
        test_metrics[name] = metrics
        res_str = f"{name}\n\n" + "\n".join([f"{k.capitalize()}: {v:.4f}" for k, v in metrics.items()]) + f"\n\n{classification_report(y_test, y_pred, target_names=['no_FHB', 'FHB'], zero_division=0)}"
        (out_dir / f"test_results_{name}.txt").write_text(res_str)
        summary_lines.append(f"{name:<50} " + " ".join([f"{k}={v:.4f}" for k, v in metrics.items()]))
        
        fig, ax = plt.subplots(figsize=(5, 4))
        sns.heatmap(confusion_matrix(y_test, y_pred), annot=True, fmt="d", cmap="Blues", xticklabels=["no_FHB", "FHB"], yticklabels=["no_FHB", "FHB"], ax=ax)
        ax.set(xlabel="Predicted", ylabel="True", title=name)
        fig.tight_layout()
        fig.savefig(out_dir / f"confusion_matrix_{name}.png", dpi=150)
        plt.close(fig)
        # joblib.dump(clf, Path(path_model) / f"{name}.joblib")
        
    (out_dir / "all_results.txt").write_text("\n".join(summary_lines))
    return test_metrics

def main(args):
    seed, iteration = getattr(args, "seed", SEED), getattr(args, "iteration", 0)
    random.seed(seed); np.random.seed(seed)
    out_dir, path_models = Path(args.out_dir) / f"iter{iteration}", Path(args.path_models) / f"iter{iteration}"
    
    X_full_train, y_full_train = load_data(PATH_DATA, "train")
    X_test, y_test = load_data(PATH_DATA, "test")
    columns = list(X_full_train.keys())
    
    cluster_models = fit_cluster_models(X_full_train, columns, args.n_clusters, seed)
    X_full_train_cl = apply_clusters(cluster_models, X_full_train, len(y_full_train))
    X_test_cl = apply_clusters(cluster_models, X_test, len(y_test))
    
    X_train_cl, X_val_cl, y_train, y_val = train_test_split(X_full_train_cl, y_full_train, test_size=args.val_size, stratify=y_full_train, random_state=seed)
    best_params = {n: study_hyperparameters(cls, PARAM_GRID[n], X_train_cl, y_train, X_val_cl, y_val) for n, cls in CLASSIFIERS.items()}
    
    X_final_train = np.concatenate([X_train_cl, X_val_cl], axis=0)
    y_final_train = np.concatenate([y_train, y_val], axis=0)
    
    best_models = {n: CLASSIFIERS[n](**best_params[n]).fit(X_final_train, y_final_train) for n in CLASSIFIERS}
    bagging_models = {n: clf.fit(X_final_train, y_final_train) for n, clf in build_bagging_classifiers(best_params, seed).items()}
    final_voting_models = {}
    if args.voting != "none":
        voting_configs = search_voting_combinations(best_params, X_train_cl, y_train, X_val_cl, y_val, min_size=args.vote_min, max_size=args.vote_max, voting=args.voting, top_k=args.vote_top_k)
        final_voting_models = {n: VotingClassifier(estimators=cfg["estimators"], voting=cfg["strategy"], n_jobs=-1).fit(X_final_train, y_final_train) for n, cfg in voting_configs.items()}
    
    all_models = {**best_models, **bagging_models, **final_voting_models}
    cv_results = cross_validate_models(all_models, X_train_cl, y_train, n_folds=args.cv_folds, seed=seed)
    test_metrics = evaluate_and_save(all_models, X_test_cl, y_test, str(path_models), str(out_dir))
    
    save_cv_results(cv_results, out_dir, test_metrics, best_params)
    print_consistency_report(cv_results, test_metrics, out_dir)
    for m in CV_SCORING: plot_cv_results(cv_results, test_metrics, out_dir, metric=m)
    
    ts = datetime.now().isoformat(timespec="seconds")
    rows = []
    for n in all_models:
        row = {"timestamp": ts, "iteration": iteration, "seed": seed, "model": n}
        row.update({f"test_{m}": v for m, v in test_metrics.get(n, {}).items()})
        if cv_results.get(n):
            row.update({f"cv_{m}_{k}": stat[k] for m, stat in cv_results[n].items() for k in ["mean", "std"]})
        rows.append(row)
        append_metrics_row(METRICS_CSV, row)
    return pd.DataFrame(rows)

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Viability classification via time-series clustering + ML")
    parser.add_argument("--n_clusters", type=int, default=N_CLUSTERS)
    parser.add_argument("--path_models", default=PATH_MODELS)
    parser.add_argument("--out_dir", default=PATH_OUTPUT)
    parser.add_argument("--vote_min", type=int, default=2)
    parser.add_argument("--vote_max", type=int, default=None)
    parser.add_argument("--voting", default="both", choices=["hard", "soft", "both", "none"])
    parser.add_argument("--vote_top_k", type=int, default=3)
    parser.add_argument("--cv_folds", type=int, default=CV_FOLDS)
    parser.add_argument("--val_size", type=float, default=VAL_SIZE)
    parser.add_argument("--iteration", type=int, default=0)
    parser.add_argument("--seed", type=int, default=SEED)
    main(parser.parse_args())